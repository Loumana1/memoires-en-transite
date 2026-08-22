"""Type ambiance MUSICALE / TEXTURE pour les pools Proto 08.

Priorité : feuille Ambiances du classeur → heuristique registre (start ≥ 1400 s).
"""
from __future__ import annotations

import csv
import os
import re

SONS_AMB = "SONS_V3/AMBIANCE"
MUSICALE = "MUSICALE"
TEXTURE = "TEXTURE"
XLSX = os.path.join("docs", "Matiere", "catalogue_fragments.xlsx")
REGISTRE = os.path.join("docs", "Matiere", "registre_ids.csv")
# Provisoire jusqu'à remplissage `type_ambiance` à l'oreille (Attributs §3).
MUSICALE_START_S = 1400.0
MIN_MUSICALE_S = 20.0
MIN_TEXTURE_S = 15.0

# Sélection Loumana — seules ambiances Cortex autorisées (Proto 08).
CURATED_AMBIANCE_STEMS = (
    "A41", "A43", "A38", "A44", "A45", "A46", "A48", "A49",
    "A53", "A55", "A59", "A61", "A62", "A64", "A69",
)
# Alias historique
CURATED_MUSICALE_STEMS = CURATED_AMBIANCE_STEMS


def _norm_type(val: str | None) -> str | None:
    if not val:
        return None
    u = str(val).strip().upper()
    if u.startswith("MUS"):
        return MUSICALE
    if u.startswith("TEX"):
        return TEXTURE
    return None


def _read_xlsx_types(root: str) -> dict[str, str]:
    path = os.path.join(root, XLSX)
    if not os.path.isfile(path):
        return {}
    try:
        import openpyxl
    except ImportError:
        return {}
    out: dict[str, str] = {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Ambiances" not in wb.sheetnames:
            return {}
        ws = wb["Ambiances"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        if not rows:
            return {}
        header = [str(c).strip().lower() if c else "" for c in rows[0]]
        if header[0] != "id":
            return {}
        try:
            ti = header.index("type ambiance")
        except ValueError:
            return {}
        for row in rows[1:]:
            if not row or not row[0]:
                continue
            sid = str(row[0]).strip()
            t = _norm_type(row[ti] if ti < len(row) else None)
            if t:
                out[sid] = t
    finally:
        wb.close()
    return out


def _read_registre_amb(root: str) -> list[dict]:
    path = os.path.join(root, REGISTRE)
    if not os.path.isfile(path):
        return []
    rows = []
    with open(path, newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row.get("role") != "AMBIANCE":
                continue
            try:
                start = float(row.get("start") or 0)
                dur = float(row.get("dur") or 0)
            except ValueError:
                continue
            sid = (row.get("id") or "").strip()
            if not sid:
                continue
            rows.append(dict(id=sid, start=start, dur=dur))
    return rows


def classify_ambiance(sid: str, start: float, manual: dict[str, str]) -> str:
    t = _norm_type(manual.get(sid))
    if t:
        return t
    if start >= MUSICALE_START_S:
        return MUSICALE
    return TEXTURE


def _rel_wav(root: str, sid: str) -> str | None:
    path = os.path.join(root, SONS_AMB, f"{sid}.wav")
    if os.path.isfile(path):
        return os.path.relpath(path, root)
    # id sans suffixe _ambiance
    alt = re.sub(r"_ambiance$", "", sid, flags=re.I)
    path2 = os.path.join(root, SONS_AMB, f"{alt}.wav")
    if os.path.isfile(path2):
        return os.path.relpath(path2, root)
    globbed = [
        p for p in os.listdir(os.path.join(root, SONS_AMB))
        if p.lower().startswith(sid.lower()[:4]) and p.lower().endswith(".wav")
    ]
    if len(globbed) == 1:
        return os.path.join(SONS_AMB, globbed[0])
    return None


def _curated_pool(by_id: dict[str, tuple[str, float]]) -> list[str]:
    """Liste exclusive Loumana — pas de filtre durée, pas de fallback."""
    out: list[str] = []
    for stem in CURATED_AMBIANCE_STEMS:
        for sid, (rel, _dur) in by_id.items():
            if sid.startswith(stem + "_"):
                out.append(rel)
                break
    return out


def ambiance_paths_by_type(root: str, verbose: bool = True) -> tuple[list[str], list[str]]:
    """(musicale, texture) — slots 32 / 33, pool curated puis séparation heuristique."""
    by_id: dict[str, tuple[str, float]] = {}
    manual = _read_xlsx_types(root)
    for row in _read_registre_amb(root):
        rel = _rel_wav(root, row["id"])
        if not rel:
            continue
        by_id[row["id"]] = (rel, row["dur"])

    pool = _curated_pool(by_id)
    if not pool:
        raise SystemExit(
            "SONS_V3/AMBIANCE: aucun fichier curated (A38, A41, A43…) — vérifier la matière"
        )

    mus: list[str] = []
    tex: list[str] = []
    for rel in pool:
        stem = os.path.splitext(os.path.basename(rel))[0]
        sid = stem.split("_")[0] if "_" in stem else stem
        row = next((r for r in _read_registre_amb(root) if r["id"].startswith(sid)), None)
        start = row["start"] if row else 0.0
        kind = classify_ambiance(sid, start, manual)
        if kind == MUSICALE:
            mus.append(rel)
        else:
            tex.append(rel)
    if not mus:
        mus = list(pool)
    if not tex:
        tex = [p for p in pool if p not in mus] or list(pool)
    # Texture : au moins 5 fichiers pour éviter la répétition sur 40 s de Cortex
    if len(tex) < 5:
        extra = [p for p in mus if p not in tex]
        tex.extend(extra[: max(0, 5 - len(tex))])
    if verbose:
        print(f"  Ambiance curated: {len(pool)} · musicale {len(mus)} · texture {len(tex)}")
    return mus, tex
