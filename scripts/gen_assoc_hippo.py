#!/usr/bin/env python3
"""Précalcule les séquences associatives Hippocampe V1 pour Pure Data.

Lit catalogue_fragments.xlsx (feuille Hippocampe) si disponible ; sinon pools
SONS_V3 + attributs mockés. Scope MVP (Q28) : APPELER, RELIER, REPONDRE,
DISPARAITRE — silences/délais plafonnés à 2 s.

Sorties :
  pd/lib/hippo_assoc/events.txt   — séquence planifiée (~50 s d'état)
  pd/lib/hippo_assoc/pool.txt     — lignes source|comportement|… pour debug
"""
from __future__ import annotations

import csv
import os
import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "proto08"))

from proto08_lib import sons_audit08 as audit  # noqa: E402

OUT_DIR = ROOT / "pd" / "lib" / "hippo_assoc"
XLSX = ROOT / "docs" / "Matiere" / "catalogue_fragments.xlsx"
HIPPO_SHEET = "Hippocampe"
STATE_MS = 50_000

# Force associative Simon — bases provisoires (H1)
FORCE_PROBA = {1: 0.25, 2: 0.55, 3: 0.80}

# Silences selon ouverture — max 2 s (TO DO 21 août)
SILENCE_RANGES = {
    "FERME": (1000, 2000),
    "PARTIELLEMENT_SUSPENDU": (500, 1500),
    "TRES_OUVERT": (0, 1000),
}

DEFAULT_ATTR = {
    "force_associative": "2",
    "ouverture": "PARTIELLEMENT_SUSPENDU",
    "mode_lecture": "INTEGRAL",
    "delai_reponse": "0.8",
    "type_association": "",
    "famille_associative": "",
    "familles_cibles": "",
}

MOCK_INTERRUPTIBLE_RATE = 0.15


def _read_hippo_catalog() -> list[dict]:
    rows: list[dict] = []
    if not XLSX.is_file():
        return rows
    try:
        import openpyxl
    except ImportError:
        print("WARN: openpyxl absent — catalogue ignoré")
        return rows
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    try:
        if HIPPO_SHEET not in wb.sheetnames:
            return rows
        ws = wb[HIPPO_SHEET]
        raw = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        if not raw or not raw[0]:
            return rows
        headers = [str(h or "").strip().lower() for h in raw[0]]

        def col(name: str) -> int | None:
            try:
                return headers.index(name)
            except ValueError:
                return None

        idx = {k: col(k.replace("_", " ") if " " not in k else k)
               for k in (
                   "id", "type", "duree_s", "type fonctionnel", "force associative",
                   "type association", "famille associative", "familles cibles",
                   "ouverture", "mode lecture", "delai reponse",
               )}
        # normaliser clés header réelles
        key_map = {}
        for i, h in enumerate(headers):
            key_map[h.replace(" ", "_")] = i

        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row or not row[0]:
                continue
            sid = str(row[0]).strip()
            data = dict(DEFAULT_ATTR)
            data["id"] = sid
            data["type"] = str(row[key_map.get("type", 1)] or "").strip() if "type" in key_map else ""
            for src, dst in (
                ("type_fonctionnel", "type_fonctionnel"),
                ("force_associative", "force_associative"),
                ("type_association", "type_association"),
                ("famille_associative", "famille_associative"),
                ("familles_cibles", "familles_cibles"),
                ("ouverture", "ouverture"),
                ("mode_lecture", "mode_lecture"),
                ("delai_reponse", "delai_reponse"),
            ):
                ci = key_map.get(src)
                if ci is not None and ci < len(row) and row[ci]:
                    data[dst] = str(row[ci]).strip()
            rows.append(data)
    finally:
        wb.close()
    return rows


def _pool_from_disk() -> tuple[list[str], list[str]]:
    u = audit.usable_files(str(ROOT), verbose=False)
    longs = u.get((1, 1), []) or []
    shorts = u.get((1, 0), []) or []
    if not longs and not shorts:
        all_h = longs + shorts
        everything = sum(u[(si, di)] for si in range(4) for di in range(3))
        longs = everything[: max(1, len(everything) // 2)]
        shorts = everything[max(1, len(everything) // 2):] or longs
    return longs, shorts


def _stem(rel: str) -> str:
    return Path(rel).stem


def _merge_catalog_pools(catalog: list[dict]) -> tuple[list[dict], list[dict]]:
    longs, shorts = _pool_from_disk()
    by_id = {r["id"]: r for r in catalog}
    long_rows: list[dict] = []
    short_rows: list[dict] = []
    for rel in longs:
        sid = _stem(rel)
        row = dict(DEFAULT_ATTR)
        row.update(by_id.get(sid, {}))
        row["id"] = sid
        row["_path"] = rel
        row["type"] = row.get("type") or "long"
        long_rows.append(row)
    for rel in shorts:
        sid = _stem(rel)
        row = dict(DEFAULT_ATTR)
        row.update(by_id.get(sid, {}))
        row["id"] = sid
        row["_path"] = rel
        row["type"] = row.get("type") or "court"
        short_rows.append(row)
    if not long_rows:
        print("WARN: pool HIPPO long vide — fallback court")
        long_rows = short_rows[:]
    if not short_rows:
        print("WARN: pool HIPPO court vide — fallback long")
        short_rows = long_rows[:]
    return long_rows, short_rows


def _force_roll(row: dict) -> bool:
    try:
        f = int(float(str(row.get("force_associative") or "2")))
    except ValueError:
        f = 2
    p = FORCE_PROBA.get(max(1, min(3, f)), 0.55)
    return random.random() < p


def _silence_ms(row: dict) -> int:
    ouv = str(row.get("ouverture") or "PARTIELLEMENT_SUSPENDU").upper().replace(" ", "_")
    lo, hi = SILENCE_RANGES.get(ouv, SILENCE_RANGES["PARTIELLEMENT_SUSPENDU"])
    return random.randint(lo, hi)


def _delai_ms(row: dict) -> int:
    raw = str(row.get("delai_reponse") or "0.8").replace(",", ".")
    try:
        sec = float(raw.split("-")[0].strip())
    except ValueError:
        sec = 0.8
    return min(int(sec * 1000), 2000)


def _interruptible(row: dict) -> int:
    mode = str(row.get("mode_lecture") or "").upper()
    if "INTERRUPT" in mode:
        return 1
    if not XLSX.is_file() or not row.get("mode_lecture"):
        return 1 if random.random() < MOCK_INTERRUPTIBLE_RATE else 0
    return 0


def _pick_other(pool: list[dict], avoid: str) -> dict | None:
    cand = [r for r in pool if r["id"] != avoid]
    return random.choice(cand) if cand else None


def _hippo_slot(layer: int) -> int:
    """Slot playlist FSM Hippo : couche N → slot (N-1)*3 (m3 dans gen_patch08)."""
    return (layer - 1) * 3


def _slot_files(slot: int, u: dict) -> list[str]:
    si, di = slot // 3, slot % 3
    return u.get((si, di), []) or []


def _path_index(path: str, slot: int, u: dict) -> int:
    files = _slot_files(slot, u)
    if not files:
        return 0
    for i, rel in enumerate(files):
        if rel == path:
            return i
    base = os.path.basename(path)
    for i, rel in enumerate(files):
        if os.path.basename(rel) == base:
            return i
    return 0


def _pick_other_multi(pool: list[dict], avoid_ids) -> dict | None:
    """Tire dans pool en évitant tous les IDs actifs (anti-doublon multi-couche)."""
    avoid = set(avoid_ids)
    cand = [r for r in pool if r["id"] not in avoid]
    return random.choice(cand) if cand else None


def build_sequence(longs: list[dict], shorts: list[dict], rng: random.Random) -> list[dict]:
    """Construit une séquence d'événements Pd pour ~50 s.

    Invariant anti-doublon : à tout instant, les sources jouées sur L1–L4
    sont deux à deux distinctes. active_sources[layer] = id en cours.

    Chaque play embarque slot + index playlist (résolu en Python) pour que
    player_state_08 ouvre la bonne ligne sans tirage aléatoire.
    """
    u = audit.usable_files(str(ROOT), verbose=False)
    events: list[dict] = []
    t = 0

    # --- Démarrage : L1 != L2 garanti ---
    long_a = rng.choice(longs)
    long_b = _pick_other_multi(longs, [long_a["id"]]) or long_a
    short_a = rng.choice(shorts)
    short_b = _pick_other_multi(shorts, [short_a["id"]]) or short_a

    # Suivi des sources actives par couche (anti-doublon continu)
    active_sources: dict[int, str] = {}

    def _play_ev(layer: int, row: dict, **extra) -> dict:
        """Construit un event play et met à jour active_sources."""
        active_sources[layer] = row["id"]
        path = row.get("_path") or "0"
        slot = _hippo_slot(layer)
        idx = _path_index(path, slot, u) if path != "0" else -1
        return {
            "t": t, "action": "play", "layer": layer,
            "source": row["id"], "slot": slot, "index": idx,
            "_path": path, **extra,
        }

    def _safe_pick(pool: list[dict]) -> dict:
        """Tire dans pool en évitant tous les actifs courants."""
        r = _pick_other_multi(pool, active_sources.values())
        return r if r is not None else rng.choice(pool)

    # L1/L2 permanents — démarrage immédiat (FSM)
    events.append(_play_ev(1, long_a))
    events.append(_play_ev(2, long_b))

    # RELIER succession : L2 après silence post L1 (simulé par délai)
    t += _silence_ms(long_a)
    long_b2 = _pick_other_multi(longs, active_sources.values()) or long_b
    events.append(_play_ev(2, long_b2,
                           comportement="RELIER", variante="succession",
                           cible=long_b2["id"]))

    # APPELER direct ou sans réponse depuis L1
    t += rng.randint(800, 2000)
    if _force_roll(long_a):
        cible = _safe_pick(shorts)
        if rng.random() < 0.7:
            events.append(_play_ev(3, cible,
                                   comportement="APPELER", variante="direct",
                                   source=long_a["id"], cible=cible["id"]))
            t += _delai_ms(long_a) + _silence_ms(cible)
        else:
            events.append({
                "t": t,
                "action": "silence",
                "comportement": "APPELER", "variante": "sans_reponse",
                "source": long_a["id"],
            })
            t += _silence_ms(long_a)

    # Courts — interférences L3/L4
    t = max(t, 1800)
    src3 = _safe_pick(shorts)
    events.append(_play_ev(3, src3,
                           comportement="REPONDRE", variante="immediate",
                           cible=src3["id"], spatial="voisin"))

    t += rng.randint(1200, 2000)
    src4 = _safe_pick(shorts)
    events.append(_play_ev(4, src4,
                           comportement="REPONDRE", variante="spatiale",
                           cible=src4["id"], spatial="oppose"))

    # DISPARAITRE nette sur L3
    t += rng.randint(600, 1500)
    fade = rng.randint(35, 100)
    active_sources.pop(3, None)
    events.append({
        "t": t, "action": "cut", "layer": 3, "fade_ms": fade,
        "comportement": "DISPARAITRE", "variante": "nette",
        "source": src3["id"],
    })

    # Motion spatiale mid-passage
    t += rng.randint(1000, 2000)
    events.append({"t": t, "action": "motion", "recipe": rng.randint(0, 4)})

    # Densité Hippo — interférences L3/L4 jusqu'à 50 s
    while t < STATE_MS - 2500:
        t += rng.randint(3500, 6500)
        if t >= STATE_MS - 1500:
            break
        layer = 3 if rng.random() < 0.55 else 4
        src = _safe_pick(shorts)
        events.append(_play_ev(layer, src,
                               comportement="REPONDRE",
                               variante="spatiale" if layer == 4 else "immediate",
                               cible=src["id"]))
        if rng.random() < 0.35:
            t += rng.randint(800, 1800)
            active_sources.pop(layer, None)
            events.append({
                "t": t, "action": "cut", "layer": layer,
                "fade_ms": rng.randint(40, 90),
                "comportement": "DISPARAITRE", "variante": "nette",
                "source": src["id"],
            })
        if rng.random() < 0.25:
            t += rng.randint(500, 1200)
            events.append({"t": t, "action": "motion", "recipe": rng.randint(0, 4)})

    # Filtrer > STATE_MS
    events = [e for e in events if e["t"] <= STATE_MS]
    events.sort(key=lambda e: e["t"])
    return events


def write_outputs(events: list[dict], pool_lines: list[str]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ev_path = OUT_DIR / "events.txt"
    with ev_path.open("w", encoding="utf-8") as fh:
        fh.write("# delay_ms action layer [fade_ms|recipe|slot index path]\n")
        fh.write("# play : slot+index = ligne playlist (anti-doublon, résolu en Python)\n")
        for e in events:
            act = e["action"]
            if act == "silence" or act == "noop":
                fh.write(f"{e['t']} {act} 0 -\n")
            elif act == "motion":
                fh.write(f"{e['t']} motion 0 {e['recipe']}\n")
            elif act == "cut":
                fh.write(f"{e['t']} cut {e['layer']} {e.get('fade_ms', 50)}\n")
            else:
                slot = e.get("slot", _hippo_slot(int(e["layer"])))
                idx = e.get("index", -1)
                path = e.get("_path") or "0"
                fh.write(f"{e['t']} play {e['layer']} {slot} {idx} {path}\n")
    pool_path = OUT_DIR / "pool.txt"
    with pool_path.open("w", encoding="utf-8") as fh:
        fh.write("# source | comportement | variante | cible | layer\n")
        for line in pool_lines:
            fh.write(line + "\n")
    print(f"OK {ev_path.relative_to(ROOT)} ({len(events)} events)")


def main() -> None:
    seed = int(os.environ.get("HIPPO_ASSOC_SEED", "42"))
    rng = random.Random(seed)
    catalog = _read_hippo_catalog()
    longs, shorts = _merge_catalog_pools(catalog)
    events = build_sequence(longs, shorts, rng)

    pool_lines = []
    for e in events:
        if e.get("comportement"):
            pool_lines.append(
                f"{e.get('source', '-')} | {e['comportement']} | {e.get('variante', '-')} "
                f"| {e.get('cible', '-')} | {e.get('layer', 0)}"
            )
    write_outputs(events, pool_lines)


if __name__ == "__main__":
    main()
